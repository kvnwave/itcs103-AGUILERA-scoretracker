from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import messagebox, scrolledtext

def validate_inputs():
    studentname = studentname_entry.get()
    score = score_entry.get()

    if not studentname or not score:
        messagebox.showerror("Error", "All fields are required.")
        return False
    
    try:
        score_value =int(score)
        if score_value < 0 or score_value > 100:
            messagebox.showerror("Score must be between 0 and 100.")
            return False
    except ValueError:
        messagebox.showerror("Score must be a numeric value.")
        return False

    return True

def save_to_excel():
    if not validate_inputs():
        return
    
    studentname = studentname_entry.get()
    score = int(score_entry.get())
    grade = grade_remarks(score)

    try:
        wb = load_workbook("student_scores.xlsx")
        ws = wb["ScoreTracker"]
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "ScoreTracker"
        ws.append(["Student Name", "Score", "Grade"])

    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == studentname:
            ws.cell(row=row, column=2, value=score)
            ws.cell(row=row, column=3, value=grade)
            messagebox.showinfo("Success", f"Updated record for {studentname}.")
            wb.save("student_scores.xlsx")
            studentname_entry.delete(0, tk.END)
            score_entry.delete(0, tk.END)
            return

    ws.append([studentname, score, grade])
    wb.save("student_scores.xlsx")
    messagebox.showinfo("Succes","Data saved to Excel.")
    studentname_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

def grade_remarks(score):
    if score <= 74:
        return "Failed"
    else:
        return "Passed"

def show_all_records():
    try:
        wb = load_workbook("student_scores.xlsx")
        ws = wb["ScoreTracker"]
    except FileNotFoundError:
        messagebox.showerror("Excel File not created yet.")
        return

    records = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=1).value
        score = ws.cell(row=row, column=2).value
        grade = ws.cell(row=row, column=3).value
        records.append(f"{name:<30} {score:<30} {grade:<10}")

    record_window = tk.Toplevel()
    record_window.title("Score Tracker")
    record_window.geometry("400x250")
    text_area = scrolledtext.ScrolledText(record_window, width=50, height=15, font=("Arial", 10))
    text_area.pack(padx=15, pady=10)
    text_area.insert(tk.END, f"{'Student Name':<30} {'Score':<30} {'Grade':<10}\n")
    text_area.insert(tk.END, "="*40 + "\n")
    for record in records:
        text_area.insert(tk.END, record + "\n")
    text_area.config(state=tk.DISABLED)

window = tk.Tk()
window.title("Score Tracker")

tk.Label(window, text="Student Name").grid(row=2, column=0, padx=20,pady=2,sticky="NSEW")
tk.Label(window, text="Score").grid(row=3, column=0,padx=20, pady=2,sticky="NSEW")

studentname_entry = tk.Entry(window)
score_entry = tk.Entry(window)

studentname_entry.grid(row=2, column=1,sticky="NSEW", pady=2)
score_entry.grid(row=3, column=1,sticky="NSEW", pady=2)

tk.Button(window, text="Submit", command=save_to_excel, bg="#1899D6", fg="#FFFFFF").grid(row=9, column=0,pady=6, ipadx=20,ipady=1, sticky="W", padx=10)
tk.Button(window, text="Show All Records", command=show_all_records, bg="#1899D6", fg="#FFFFFF").grid(row=9,column=1, pady=6, sticky="W",padx=10, ipady=2)

window.mainloop()